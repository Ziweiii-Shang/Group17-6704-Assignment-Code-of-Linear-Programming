"""Recompute Table 2 and the network statistics discussed in Sections 6.1/6.3.

Mean path length counts arcs on a minimum-travel-time path, breaking ties by
fewer arcs. It is not the minimum-hop distance. All commodities have equal
weight, including zero-hop same-origin/destination commodities.
"""
import argparse
import csv
import heapq
from pathlib import Path
import pickle
import statistics
import sys
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts/pte'))
from run_pte_experiments import bounds, graph_data


def fastest_path_hops(outgoing, origin, destination):
    labels = [(10**18, 10**9)] * len(outgoing)
    labels[origin] = (0, 0)
    heap = [(0, 0, origin)]
    while heap:
        elapsed, hops, node = heapq.heappop(heap)
        if (elapsed, hops) != labels[node]:
            continue
        if node == destination:
            return hops
        for _, neighbor, duration in outgoing[node]:
            candidate = (elapsed + duration, hops + 1)
            if candidate < labels[neighbor]:
                labels[neighbor] = candidate
                heapq.heappush(heap, (*candidate, neighbor))
    raise ValueError('Unreachable commodity destination')


def reachable(graph):
    seen, pending = {0}, [0]
    while pending:
        for _, neighbor, _ in graph[pending.pop()]:
            if neighbor not in seen:
                seen.add(neighbor); pending.append(neighbor)
    return len(seen) == len(graph)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', type=Path, default=ROOT / 'data/processed')
    parser.add_argument('--raw', type=Path, default=ROOT / 'data/raw')
    parser.add_argument('--output', type=Path, default=ROOT / 'results/report/network_statistics.csv')
    args = parser.parse_args()
    rows = []
    for family in ('corridor', 'hub', 'random'):
        for size in ('small', 'middle', 'Large'):
            name = f'instance_{family}_timewindow_{size}'
            data = pickle.loads((args.input / f'{name}.input.pkl').read_bytes())
            n, m, nk = len(data['nodes']), len(data['arcs']), len(data['commodities'])
            outgoing, distances, totals, _, _ = bounds(data, range(nk))
            _, reverse = graph_data(data)
            book = load_workbook(args.raw / f'{name}.xlsx', read_only=True, data_only=True)
            zones = {}
            for _, i, j, _, _, _, zi, zj in book['Arcs'].iter_rows(min_row=2, values_only=True):
                zones[i], zones[j] = zi, zj
            commodity_rows = list(book['Commodities'].iter_rows(min_row=2, values_only=True))
            cross_region = sum(row[4] == 'cross_region' for row in commodity_rows)
            cross_zone = sum(zones[row[1]] != zones[row[2]] for row in commodity_rows)
            book.close()
            hops, allowed = [], 0
            for k, (_, origin, destination, _, budget) in enumerate(data['commodities']):
                hops.append(fastest_path_hops(outgoing, origin, destination))
                forward, backward = distances[k]
                allowed += sum(forward[i] + tau + backward[j] <= budget for _, i, j, _, tau in data['arcs'])
            rows.append({'instance': name, 'family': family, 'size': size, 'nodes': n, 'arcs': m,
                         'commodities': nk, 'compact_variables': nk * m, 'compact_rows': nk * n + m,
                         'mean_fastest_path_arcs': statistics.mean(hops),
                         'time_feasible_commodity_arc_pairs': allowed,
                         'time_feasible_commodity_arc_pct': 100 * allowed / (nk * m),
                         'strongly_connected': reachable(outgoing) and reachable(reverse),
                         'configured_cross_region_od_ratio': data['config']['CROSS_REGION_OD_RATIO'],
                         'observed_cross_region_od_pct': 100 * cross_region / nk,
                         'observed_cross_zone_endpoint_pct': 100 * cross_zone / nk,
                         'full_variables': totals['full_transport'] + totals['full_sink'],
                         'bound_variables': totals['bound_transport'] + totals['bound_sink'],
                         'source_sha256': data['source_sha256']})
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open('w', newline='', encoding='utf-8') as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0])); writer.writeheader(); writer.writerows(rows)
    print(f'Wrote {len(rows)} instance statistics to {args.output}')


if __name__ == '__main__':
    main()
