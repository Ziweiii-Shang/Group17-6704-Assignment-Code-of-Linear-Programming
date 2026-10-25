"""Build Tables 5--7 and Figures 3--5 from recorded or freshly run results.

Default inputs are the submitted report's historical measurements. To analyse
fresh runs, explicitly supply all three result directories; missing instances
are errors, never silently filled with old measurements.
"""
import argparse
import csv
import json
import math
import os
from pathlib import Path
import statistics
import tempfile

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
FAMILIES = ('corridor', 'hub', 'random')
SIZES = ('small', 'middle', 'Large')
INSTANCES = [f'instance_{f}_timewindow_{s}' for f in FAMILIES for s in SIZES]
# Explicitly reconciled against Group17.pdf Table 6. The source workbook's
# numeric suffix legend is inconsistent with its rows; do not infer names.
SOURCE_ROWS = {f'{code}1-{number}': f'instance_{family}_timewindow_{size}'
               for code, family in zip(('L', 'M', 'S'), FAMILIES)
               for number, size in zip((1, 2, 3), ('Large', 'middle', 'small'))}


def source_name(path):
    try:
        return path.resolve().relative_to(ROOT).as_posix()
    except ValueError:
        return path.name


def write_csv(path, rows):
    with path.open('w', newline='', encoding='utf-8') as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def read_historical_cg(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows = {}
    for source_row, values in enumerate(workbook['Sheet1'].iter_rows(min_row=2, max_row=10, values_only=True), 2):
        label, cost, columns, seconds = values
        instance = SOURCE_ROWS[label]
        rows[instance] = {'instance': instance, 'cg_objective': float(cost),
                          'cg_seconds': float(str(seconds).removesuffix('s').strip()),
                          'cg_columns': int(columns), 'source_label': label,
                          'source': f'{path.name}:Sheet1!A{source_row}:D{source_row}',
                          'source_sha256': __import__('hashlib').sha256(path.read_bytes()).hexdigest()}
    workbook.close()
    return rows


def read_fresh_cg(directory):
    rows = {}
    for instance in INSTANCES:
        path = directory / f'{instance}.xlsx'
        workbook = load_workbook(path, read_only=True, data_only=True)
        metrics = dict(workbook['Summary'].iter_rows(min_row=2, values_only=True))
        history = list(workbook['Iterations'].values)
        last = dict(zip(history[0], history[-1]))
        if metrics['final_artificial_flow'] > 1e-8 or last['min_reduced_cost'] < -1e-8:
            raise ValueError(f'{instance}: CG has not certified convergence with zero artificial flow')
        if Path(metrics['instance_file']).stem != instance:
            raise ValueError(f'{instance}: result workbook identifies a different input')
        rows[instance] = {'instance': instance, 'cg_objective': metrics['final_transport_objective'],
                          'cg_seconds': metrics['total_runtime_sec'],
                          'cg_columns': metrics['generated_columns'], 'source_label': instance,
                          'source': f'{path.name}:Summary',
                          'source_sha256': __import__('hashlib').sha256(path.read_bytes()).hexdigest()}
        workbook.close()
    return rows


def read_heuristics(directory):
    baseline = json.loads((directory / 'results.json').read_text())
    retuned = json.loads((directory / 'results_ga_retuned.json').read_text())
    rows = {(r['instance'], r['method']): r for r in baseline}
    if len(rows) != 18 or len(baseline) != 18:
        raise ValueError('Expected nine baseline instances for each of two methods')
    expected = {('instance_corridor_timewindow_middle', 'genetic'),
                ('instance_corridor_timewindow_Large', 'genetic')}
    if {(r['instance'], r['method']) for r in retuned} != expected or len(retuned) != 2:
        raise ValueError('Expected the two original Corridor GA retunings')
    for row in retuned:
        rows[row['instance'], row['method']] = row
    for instance in INSTANCES:
        for method in ('k-neighborhood', 'genetic'):
            row = rows[instance, method]
            if not row['feasible'] or row['objective'] is None or row['demand_short'] > 1e-6:
                raise ValueError(f'{instance}/{method}: final report requires a feasible result')
    return rows


def plots(directory, cg, comparison):
    os.environ.setdefault('MPLCONFIGDIR', str(Path(tempfile.gettempdir()) / 'group17-mpl-cache'))
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    plt.rcParams.update({'font.size': 10, 'axes.spines.top': False, 'axes.spines.right': False})
    labels = ['Small', 'Medium', 'Large']
    fig, ax = plt.subplots(figsize=(6, 4))
    for family in FAMILIES:
        ax.plot(labels, [cg[f'instance_{family}_timewindow_{s}']['cg_seconds'] for s in SIZES],
                marker='o', label=family.title())
    ax.set(xlabel='Instance size', ylabel='CG recorded runtime (s)')
    ax.legend(); fig.tight_layout(); fig.savefig(directory / 'figure3_cg_runtime.png', dpi=180); plt.close(fig)
    fig, ax = plt.subplots(figsize=(10, 4))
    x = list(range(9))
    ax.bar([i - .2 for i in x], [r['kn_gap_cg_pct'] for r in comparison], .4, label='K-neighborhood')
    ax.bar([i + .2 for i in x], [r['ga_gap_cg_pct'] for r in comparison], .4, label='GA (updated budgets)')
    ax.set_xticks(x, [f'{r["family"].title()} / {r["size"]}' for r in comparison], rotation=30, ha='right')
    ax.set_ylabel('Cost gap relative to CG (%)'); ax.legend()
    fig.tight_layout(); fig.savefig(directory / 'figure4_heuristic_gaps.png', dpi=180); plt.close(fig)
    fig, axes = plt.subplots(1, 3, figsize=(11, 4), sharey=True)
    for ax, family in zip(axes, FAMILIES):
        group = [r for r in comparison if r['family'] == family]
        for key, label in [('cg_seconds', 'CG'), ('kn_search_seconds', 'K-neighborhood'), ('ga_search_seconds', 'GA')]:
            ax.plot(labels, [r[key] for r in group], marker='o', label=label)
        ax.set(title=family.title(), xlabel='Size', yscale='log')
    axes[0].set_ylabel('Recorded runtime (s; log scale)'); axes[-1].legend()
    fig.text(.5, .01, 'Heuristic search excludes path-pool construction; these are recorded, not controlled speedup comparisons.',
             ha='center', fontsize=8)
    fig.tight_layout(rect=(0, .04, 1, 1)); fig.savefig(directory / 'figure5_recorded_times.png', dpi=180); plt.close(fig)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--cg-workbook', type=Path, default=ROOT / 'results/cg_original.xlsx')
    parser.add_argument('--cg-dir', type=Path, help='Nine newly generated CG workbooks; replaces historical workbook')
    parser.add_argument('--heuristic-dir', type=Path)
    parser.add_argument('--pte-dir', type=Path)
    parser.add_argument('--output', type=Path, default=ROOT / 'results/report')
    args = parser.parse_args()
    if any((args.cg_dir, args.heuristic_dir, args.pte_dir)) and not all((args.cg_dir, args.heuristic_dir, args.pte_dir)):
        parser.error('For fresh comparisons, supply --cg-dir, --heuristic-dir and --pte-dir together.')
    args.heuristic_dir = args.heuristic_dir or ROOT / 'results'
    args.pte_dir = args.pte_dir or ROOT / 'results'
    cg = read_fresh_cg(args.cg_dir) if args.cg_dir else read_historical_cg(args.cg_workbook)
    heuristics = read_heuristics(args.heuristic_dir)
    with (args.pte_dir / 'main_results.csv').open(encoding='utf-8-sig') as handle:
        pte = {r['instance']: r for r in csv.DictReader(handle)}
    if set(cg) != set(INSTANCES) or set(pte) != set(INSTANCES):
        raise ValueError('Exactly nine named instances are required')
    comparison, random = [], []
    for instance in INSTANCES:
        k = heuristics[instance, 'k-neighborhood']; g = heuristics[instance, 'genetic']; c = cg[instance]
        family, size = instance.split('_')[1], instance.split('_')[-1]
        row = {'instance': instance, 'family': family, 'size': size,
               'cg_objective': c['cg_objective'], 'cg_seconds': c['cg_seconds'], 'cg_columns': c['cg_columns'],
               'kn_objective': k['objective'], 'kn_gap_cg_pct': 100 * (k['objective'] / c['cg_objective'] - 1),
               'kn_search_seconds': k['seconds'], 'kn_pool_seconds': k['pool_seconds'],
               'ga_objective': g['objective'], 'ga_gap_cg_pct': 100 * (g['objective'] / c['cg_objective'] - 1),
               'ga_width': g['detail']['width'], 'ga_search_seconds': g['seconds'], 'ga_pool_seconds': g['pool_seconds'],
               'pte_status': pte[instance]['main_status'], 'pte_objective': pte[instance]['objective'],
               'pte_total_seconds': pte[instance]['total_seconds']}
        comparison.append(row)
        if family == 'random':
            if row['pte_status'] != 'Optimal' or not row['pte_objective']:
                raise ValueError(f'{instance}: PTE needs an optimal result for Table 5')
            random.append({'instance': instance, 'pte_objective': float(row['pte_objective']),
                           'cg_objective': c['cg_objective'], 'pte_seconds': float(row['pte_total_seconds']),
                           'cg_seconds': c['cg_seconds'],
                           'absolute_objective_difference': abs(float(row['pte_objective'])-c['cg_objective']),
                           'relative_objective_difference_pct': 100*abs(float(row['pte_objective'])-c['cg_objective'])/max(1,abs(c['cg_objective']))})
    for row in comparison:
        if not all(math.isfinite(float(row[key])) for key in ['cg_objective','cg_seconds','kn_objective','ga_objective']):
            raise ValueError(f'Non-finite result: {row["instance"]}')
    summary = {'instances': 9, 'mean_kn_gap_cg_pct': statistics.mean(r['kn_gap_cg_pct'] for r in comparison),
               'mean_ga_gap_cg_pct': statistics.mean(r['ga_gap_cg_pct'] for r in comparison),
               'mean_kn_recorded_time_reduction_pct': statistics.mean(100 * (1 - r['kn_search_seconds'] / r['cg_seconds']) for r in comparison),
               'mean_kn_time_reduction_using_table_rounded_times_pct': statistics.mean(100 * (1 - round(r['kn_search_seconds'], 2) / round(r['cg_seconds'], 2)) for r in comparison),
               'kn_runtime_scope': 'search only; excludes pool construction',
               'ga_runtime_scope': 'search across all three seeds; excludes pool construction',
               'comparison_scope': 'recorded runtimes, not a same-machine controlled speedup',
               'cg_input': source_name(args.cg_dir or args.cg_workbook),
               'heuristic_input': source_name(args.heuristic_dir), 'pte_input': source_name(args.pte_dir)}
    args.output.mkdir(parents=True, exist_ok=True)
    write_csv(args.output / 'cg_results.csv', [cg[k] for k in INSTANCES])
    write_csv(args.output / 'pte_cg_random.csv', random)
    write_csv(args.output / 'method_comparison.csv', comparison)
    model_sizes = []
    with (args.pte_dir / 'main_results.csv').open(encoding='utf-8-sig') as handle:
        for record in csv.DictReader(handle):
            full, pte_variables = int(record['full_variables']), int(record['pte_variables'])
            model_sizes.append({'instance': record['instance'], 'full_te_variables': full,
                                'retained_pte_variables': pte_variables,
                                'variables_removed_pct': 100 * (1 - pte_variables / full),
                                'pte_status': record['main_status']})
    if {row['instance'] for row in model_sizes} != set(INSTANCES):
        raise ValueError('Full-TE/PTE size table requires all nine instance rows')
    write_csv(args.output / 'pte_model_sizes.csv', model_sizes)
    (args.output / 'comparison_summary.json').write_text(json.dumps(summary, indent=2), encoding='utf-8')
    plots(args.output, cg, comparison)
    print(f'Wrote report tables and Figures 3--5 to {args.output}')
    print(f'Mean gaps: K-neighborhood {summary["mean_kn_gap_cg_pct"]:.4f}%; GA {summary["mean_ga_gap_cg_pct"]:.4f}%')


if __name__ == '__main__':
    main()
