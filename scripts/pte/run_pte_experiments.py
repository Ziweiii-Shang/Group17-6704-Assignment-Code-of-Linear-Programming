"""Exact static PTE-LP experiments, with explicit resource limits and raw logs."""
from __future__ import annotations

import argparse
from array import array
from collections import defaultdict, deque
from decimal import Decimal
import gc
import hashlib
import heapq
import json
import os
from pathlib import Path
import pickle
import platform
import resource
import time

import highspy
import numpy as np
import psutil
from scipy.sparse import coo_matrix, csr_matrix
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
DEFAULT_INPUT = ROOT / 'data' / 'processed'
OUT = ROOT / 'results' / 'generated' / 'pte_experiments'


def millis(value):
    x = Decimal(str(value)) * 1000
    if x != x.to_integral_value():
        raise ValueError(f'Cannot scale time exactly: {value}')
    return int(x)


def read_instance(path):
    start = time.perf_counter()
    book = load_workbook(path, read_only=True, data_only=True)
    config = dict(list(book['Config'].values)[1:])
    raw_arcs = list(book['Arcs'].values)[1:]
    nodes = sorted({v for row in raw_arcs for v in row[1:3]})
    node_index = {v: i for i, v in enumerate(nodes)}
    arcs = [(row[0], node_index[row[1]], node_index[row[2]], float(row[3]), millis(row[4])) for row in raw_arcs]
    assert all(a[3] >= 0 and a[4] > 0 for a in arcs)
    arc_index = {(row[1], row[2]): i for i, row in enumerate(raw_arcs)}
    assert len(arc_index) == len(arcs), 'Costs require unambiguous physical arc endpoints'
    commodities = []
    for row in list(book['Commodities'].values)[1:]:
        budget = min(millis(row[7]), millis(row[6]) - millis(row[5]))
        assert budget >= 0 and row[3] > 0
        commodities.append((row[0], node_index[row[1]], node_index[row[2]], float(row[3]), budget))
    k_index = {c[0]: i for i, c in enumerate(commodities)}
    costs = np.full((len(commodities), len(arcs)), np.nan)
    for row in book['Costs'].iter_rows(min_row=2, values_only=True):
        k, a = k_index[row[0]], arc_index[row[1], row[2]]
        assert np.isnan(costs[k, a]), 'Duplicate cost record'
        costs[k, a] = row[3]
    book.close()
    assert np.isfinite(costs).all() and (costs >= 0).all()
    data = {'config': config, 'nodes': nodes, 'arcs': arcs, 'commodities': commodities, 'costs': costs}
    data['input_seconds'] = time.perf_counter() - start
    data['source_sha256'] = hashlib.sha256(path.read_bytes()).hexdigest()
    return data


def graph_data(data):
    n = len(data['nodes'])
    outgoing, reverse = [[] for _ in range(n)], [[] for _ in range(n)]
    for a, (_, i, j, _, tau) in enumerate(data['arcs']):
        outgoing[i].append((a, j, tau))
        reverse[j].append((a, i, tau))
    return outgoing, reverse


def dijkstra(graph, source):
    dist = [10**18] * len(graph)
    dist[source] = 0
    heap = [(0, source)]
    while heap:
        t, i = heapq.heappop(heap)
        if t != dist[i]:
            continue
        for _, j, tau in graph[i]:
            new = t + tau
            if new < dist[j]:
                dist[j] = new
                heapq.heappush(heap, (new, j))
    return dist


def bounds(data, indices):
    start = time.perf_counter()
    outgoing, reverse = graph_data(data)
    n = len(data['nodes'])
    count = defaultdict(int)
    distances = {}
    records = []
    for k in indices:
        name, o, d, q, budget = data['commodities'][k]
        f = dijkstra(outgoing, o)
        r = dijkstra(reverse, d)
        distances[k] = (f, r)
        fs = n * (budget + 1)
        fe = sum(max(0, budget - tau + 1) for _, _, _, _, tau in data['arcs'])
        bs = sum(max(0, budget - f[v] - r[v] + 1) for v in range(n))
        be = sum(max(0, budget - f[i] - tau - r[j] + 1) for _, i, j, _, tau in data['arcs'])
        gs = max(0, budget - f[d] + 1)
        vals = {'full_states': fs, 'full_transport': fe, 'full_sink': budget + 1,
                'bound_states': bs, 'bound_transport': be, 'bound_sink': gs}
        for key, value in vals.items():
            count[key] += value
        records.append({'k': k, 'commodity': name, 'budget': budget, 'earliest': f[d],
                        'same_od': o == d, **vals})
    count['individually_infeasible'] = sum(x['earliest'] > x['budget'] for x in records)
    return outgoing, distances, dict(count), records, time.perf_counter() - start


class ExpansionLimit(Exception):
    pass


def expand(data, indices, outgoing, distances, use_bounds, max_edges, max_states, seconds):
    start = time.perf_counter()
    n = len(data['nodes'])
    tails, heads, arc_ids, ks = [array('i') for _ in range(4)]
    state_nodes, state_times = array('i'), array('q')
    sinks, sources, demands, sink_ks = [], [], [], []
    commodity_records = []
    status = 'complete'
    try:
        for ordinal, k in enumerate(indices):
            name, o, d, q, budget = data['commodities'][k]
            r = distances[k][1]
            base = len(state_nodes)
            keys = array('q', [o])
            mapping = {o: base}
            state_nodes.append(o)
            state_times.append(0)
            sources.append(base)
            demands.append(q)
            before_edges = len(tails)
            local_sinks = []
            cursor = 0
            while cursor < len(keys):
                key = keys[cursor]
                t, i = divmod(key, n)
                row = base + cursor
                if i == d:
                    local_sinks.append(row)
                for a, j, tau in outgoing[i]:
                    arrival = t + tau
                    if arrival > budget or (use_bounds and arrival + r[j] > budget):
                        continue
                    newkey = arrival * n + j
                    if newkey not in mapping:
                        mapping[newkey] = len(state_nodes)
                        keys.append(newkey)
                        state_nodes.append(j)
                        state_times.append(arrival)
                    tails.append(row)
                    heads.append(mapping[newkey])
                    arc_ids.append(a)
                    ks.append(k)
                cursor += 1
                if cursor % 4096 == 0:
                    if len(tails) > max_edges or len(state_nodes) > max_states:
                        raise ExpansionLimit('representation_limit')
                    if time.perf_counter() - start > seconds:
                        raise ExpansionLimit('preprocessing_time_limit')
                    if psutil.virtual_memory().available < 3_000_000_000:
                        raise ExpansionLimit('available_memory_limit')
            sinks.extend(local_sinks)
            sink_ks.extend([ordinal] * len(local_sinks))
            commodity_records.append({'k': k, 'commodity': name, 'states': len(keys),
                                      'transport': len(tails) - before_edges, 'sink': len(local_sinks)})
            if ordinal % 50 == 0:
                print(f'  {"PTE" if use_bounds else "reachable-TE"} k={ordinal}/{len(indices)} '
                      f'states={len(state_nodes)} transport={len(tails)}', flush=True)
            if len(tails) > max_edges or len(state_nodes) > max_states:
                raise ExpansionLimit('representation_limit')
    except ExpansionLimit as exc:
        status = str(exc)
    stats = {'status': status, 'completed_commodities': len(commodity_records),
             'states': len(state_nodes), 'transport': len(tails), 'sink': len(sinks),
             'expansion_seconds': time.perf_counter() - start,
             'commodity_records': commodity_records}
    structure = dict(tails=tails, heads=heads, arc_ids=arc_ids, ks=ks,
                     state_nodes=state_nodes, state_times=state_times, sinks=sinks,
                     sources=sources, demands=demands, sink_ks=sink_ks)
    return stats, structure


def reverse_check(g):
    start = time.perf_counter()
    n = len(g['state_nodes'])
    tails, heads = np.frombuffer(g['tails'], dtype=np.int32), np.frombuffer(g['heads'], dtype=np.int32)
    incoming = coo_matrix((np.ones(len(tails), dtype=np.int8), (heads, tails)), shape=(n, n)).tocsr()
    seen = np.zeros(n, dtype=bool)
    queue = deque(g['sinks'])
    seen[g['sinks']] = True
    while queue:
        j = queue.popleft()
        for i in incoming.indices[incoming.indptr[j]:incoming.indptr[j + 1]]:
            if not seen[i]:
                seen[i] = True
                queue.append(int(i))
    return {'backward_seconds': time.perf_counter() - start,
            'backward_removed_states': int(n - seen.sum()),
            'backward_removed_transport': int((~(seen[tails] & seen[heads])).sum())}


def dense_expand(data, indices, max_edges, max_states):
    start = time.perf_counter()
    n = len(data['nodes'])
    tails, heads, arc_ids, ks = [array('i') for _ in range(4)]
    state_nodes, state_times = array('i'), array('q')
    sinks, sources, demands, sink_ks, records = [], [], [], [], []
    for ordinal, k in enumerate(indices):
        name, o, d, q, budget = data['commodities'][k]
        ns = n * (budget + 1)
        ne = sum(max(0, budget - tau + 1) for _, _, _, _, tau in data['arcs'])
        if ns + len(state_nodes) > max_states or ne + len(tails) > max_edges:
            return {'status': 'representation_limit', 'states': len(state_nodes), 'transport': len(tails),
                    'sink': len(sinks), 'completed_commodities': ordinal,
                    'expansion_seconds': time.perf_counter() - start, 'commodity_records': records}, None
        base = len(state_nodes)
        state_nodes.frombytes(np.tile(np.arange(n, dtype=np.int32), budget + 1).tobytes())
        state_times.frombytes(np.repeat(np.arange(budget + 1, dtype=np.int64), n).tobytes())
        sources.append(base + o)
        demands.append(q)
        sinks.extend((base + np.arange(budget + 1, dtype=np.int32) * n + d).tolist())
        sink_ks.extend([ordinal] * (budget + 1))
        for a, (_, i, j, _, tau) in enumerate(data['arcs']):
            copies = max(0, budget - tau + 1)
            times = np.arange(copies, dtype=np.int32)
            tails.frombytes((base + times * n + i).tobytes())
            heads.frombytes((base + (times + tau) * n + j).tobytes())
            arc_ids.frombytes(np.full(copies, a, dtype=np.int32).tobytes())
            ks.frombytes(np.full(copies, k, dtype=np.int32).tobytes())
        records.append({'k': k, 'commodity': name, 'states': ns, 'transport': ne, 'sink': budget + 1})
    stats = {'status': 'complete', 'states': len(state_nodes), 'transport': len(tails), 'sink': len(sinks),
             'completed_commodities': len(indices), 'expansion_seconds': time.perf_counter() - start,
             'commodity_records': records}
    g = dict(tails=tails, heads=heads, arc_ids=arc_ids, ks=ks, state_nodes=state_nodes,
             state_times=state_times, sinks=sinks, sources=sources, demands=demands, sink_ks=sink_ks)
    return stats, g


def optimize(data, indices, g, limit, log_path):
    start = time.perf_counter()
    ns, ne, nk = len(g['state_nodes']), len(g['tails']), len(indices)
    na, ng = len(data['arcs']), len(g['sinks'])
    nv, nr = ne + ng, ns + nk + na
    tail = np.frombuffer(g['tails'], dtype=np.int32)
    head = np.frombuffer(g['heads'], dtype=np.int32)
    aid = np.frombuffer(g['arc_ids'], dtype=np.int32)
    ki = np.frombuffer(g['ks'], dtype=np.int32)
    # Each transportation column has two balance coefficients and one capacity coefficient.
    starts = np.concatenate((np.arange(ne + 1, dtype=np.int64) * 3,
                             3 * ne + np.arange(1, ng + 1, dtype=np.int64) * 2))
    ridx = np.empty(3 * ne + 2 * ng, dtype=np.int32)
    vals = np.empty(len(ridx), dtype=np.float64)
    ridx[:3 * ne:3], ridx[1:3 * ne:3], ridx[2:3 * ne:3] = tail, head, ns + nk + aid
    vals[:3 * ne:3], vals[1:3 * ne:3], vals[2:3 * ne:3] = 1, -1, 1
    ridx[3 * ne::2], ridx[3 * ne + 1::2] = g['sinks'], ns + np.array(g['sink_ks'], dtype=np.int32)
    vals[3 * ne::2], vals[3 * ne + 1::2] = 1, -1
    b = np.zeros(ns + nk)
    b[g['sources']] = g['demands']
    b[ns:] = -np.array(g['demands'])
    caps = np.array([a[3] for a in data['arcs']])
    lp = highspy.HighsLp()
    lp.num_col_, lp.num_row_ = nv, nr
    lp.col_cost_ = np.concatenate((data['costs'][ki, aid], np.zeros(ng)))
    lp.col_lower_ = np.zeros(nv)
    lp.col_upper_ = np.full(nv, highspy.kHighsInf)
    lp.row_lower_ = np.concatenate((b, np.full(na, -highspy.kHighsInf)))
    lp.row_upper_ = np.concatenate((b, caps))
    lp.a_matrix_.format_ = highspy.MatrixFormat.kColwise
    lp.a_matrix_.start_, lp.a_matrix_.index_, lp.a_matrix_.value_ = starts, ridx, vals
    solver = highspy.Highs()
    for option, value in [('solver', 'simplex'), ('threads', 1), ('presolve', 'on'),
                          ('time_limit', limit), ('primal_feasibility_tolerance', 1e-7),
                          ('dual_feasibility_tolerance', 1e-7), ('random_seed', 42),
                          ('log_file', str(log_path)), ('log_to_console', False)]:
        assert solver.setOptionValue(option, value) == highspy.HighsStatus.kOk
    assert solver.passModel(lp) == highspy.HighsStatus.kOk
    build_seconds = time.perf_counter() - start
    begin_solve = time.perf_counter()
    solver.run()
    solve_seconds = time.perf_counter() - begin_solve
    status = solver.modelStatusToString(solver.getModelStatus())
    info = solver.getInfo()
    result = {'lp_status': status, 'variables': nv, 'constraints': nr, 'nonzeros': len(ridx),
              'build_seconds': build_seconds, 'solve_seconds': solve_seconds,
              'solver_reported_seconds': solver.getRunTime(),
              'simplex_iterations': info.simplex_iteration_count}
    if solver.getModelStatus() == highspy.HighsModelStatus.kOptimal:
        result['objective'] = solver.getObjectiveValue()
        validation_start = time.perf_counter()
        solution = solver.getSolution()
        x = np.array(solution.col_value)
        activities = np.array(solution.row_value)
        flow_activity = np.bincount(tail, weights=x[:ne], minlength=ns) - np.bincount(head, weights=x[:ne], minlength=ns)
        flow_activity += np.bincount(g['sinks'], weights=x[ne:], minlength=ns)
        sink_totals = np.bincount(g['sink_ks'], weights=x[ne:], minlength=nk)
        usage = np.bincount(aid, weights=x[:ne], minlength=na)
        times = np.frombuffer(g['state_times'], dtype=np.int64)
        durations = np.array([a[4] for a in data['arcs']], dtype=np.int64)
        nodes = np.frombuffer(g['state_nodes'], dtype=np.int32)
        budgets = np.array([c[4] for c in data['commodities']], dtype=np.int64)
        result.update(max_flow_balance_residual=float(np.max(np.abs(flow_activity - b[:ns]))),
                      max_demand_residual=float(np.max(np.abs(sink_totals - np.array(g['demands'])))),
                      max_capacity_violation=float(np.max(np.maximum(usage - caps, 0))),
                      max_nonnegativity_violation=float(max(0, -x.min(initial=0))),
                      time_transition_residual_ms=int(np.max(np.abs(times[head] - times[tail] - durations[aid]), initial=0)),
                      max_state_deadline_violation_ms=int(np.maximum(times[head] - budgets[ki], 0).max(initial=0)),
                      positive_transport_variables=int((x[:ne] > 1e-7).sum()),
                      used_physical_arcs=int((usage > 1e-7).sum()),
                      saturated_physical_arcs=int((np.abs(usage - caps) <= 1e-6).sum()),
                      duality_gap_absolute=float(abs(lp.col_cost_ @ x - b @ np.array(solution.row_dual[:ns + nk]) - caps @ np.array(solution.row_dual[ns + nk:]))),
                      validation_seconds=time.perf_counter() - validation_start)
        export_start = time.perf_counter()
        np.savez_compressed(log_path.with_suffix('.solution.npz'), x=x, physical_usage=usage)
        result['solution_write_seconds'] = time.perf_counter() - export_start
    return result


def run(path, args, cached=False, subset=None, method='pte', suffix=''):
    started = time.perf_counter()
    OUT.mkdir(parents=True, exist_ok=True)
    is_cache = path.name.endswith('.input.pkl')
    instance_name = path.name[:-len('.input.pkl')] if is_cache else path.stem
    cache = path if is_cache else ROOT / 'data' / 'processed' / f'{instance_name}.input.pkl'
    if (cached or is_cache) and cache.exists():
        with cache.open('rb') as f:
            data = pickle.load(f)
        actual_input = time.perf_counter() - started
    else:
        data = read_instance(path)
        actual_input = data['input_seconds']
        with cache.open('wb') as f:
            pickle.dump(data, f, protocol=pickle.HIGHEST_PROTOCOL)
    indices = list(range(len(data['commodities']))) if subset is None else subset
    print(instance_name, method, 'input', round(actual_input, 3), flush=True)
    outgoing, distances, counts, bound_records, bound_seconds = bounds(data, indices)
    if method == 'dense_te':
        expansion, g = dense_expand(data, indices, args.max_edges, args.max_states)
    else:
        expansion, g = expand(data, indices, outgoing, distances, method == 'pte',
                              args.max_edges, args.max_states, args.preprocess_limit)
    result = {'instance': instance_name, 'method': method, 'source_sha256': data['source_sha256'],
              'config': data['config'], 'nodes': len(data['nodes']), 'arcs': len(data['arcs']),
              'commodities': len(indices), 'commodity_indices': indices,
              'same_od_commodities': sum(data['commodities'][k][1] == data['commodities'][k][2] for k in indices),
              'input_seconds': actual_input, 'input_mode': 'pickle_cache' if cached else 'xlsx',
              'original_xlsx_input_seconds': data['input_seconds'],
              'bounds_seconds': bound_seconds, **counts, **expansion,
              'bound_records': bound_records,
              'limits': {'max_transport_edges': args.max_edges, 'max_states': args.max_states,
                         'preprocess_seconds': args.preprocess_limit, 'solve_seconds': args.solve_limit}}
    key = f'{instance_name}.{method}{suffix}'
    if expansion['status'] == 'complete':
        if method != 'dense_te':
            result.update(reverse_check(g))
        if method == 'pte':
            assert result['backward_removed_states'] == result['backward_removed_transport'] == 0
        result.update(optimize(data, indices, g, args.solve_limit, OUT / f'{key}.log'))
    result['preprocess_seconds'] = result['bounds_seconds'] + result['expansion_seconds'] + result.get('backward_seconds', 0)
    result['total_seconds'] = result['input_seconds'] + result['preprocess_seconds'] + result.get('build_seconds', 0) + result.get('solve_seconds', 0)
    result['observed_wall_seconds'] = time.perf_counter() - started
    result['peak_process_rss_bytes'] = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
    result['environment'] = {'os': platform.platform(), 'machine': platform.machine(),
                             'python': platform.python_version(), 'highs': highspy.Highs().version(),
                             'numpy': np.__version__, 'logical_cpus': psutil.cpu_count(),
                             'physical_cpus': psutil.cpu_count(logical=False),
                             'ram_bytes': psutil.virtual_memory().total}
    output_started = time.perf_counter()
    (OUT / f'{key}.json').write_text(json.dumps(result, indent=2), encoding='utf-8')
    result['json_write_seconds'] = time.perf_counter() - output_started
    (OUT / f'{key}.json').write_text(json.dumps(result, indent=2), encoding='utf-8')
    print('RESULT', key, result['status'], result.get('lp_status'), result.get('objective'),
          'states', result['states'], 'transport', result['transport'], 'time', round(result['total_seconds'], 3), flush=True)
    return result


def main():
    p = argparse.ArgumentParser()
    p.add_argument('--input', type=Path, default=DEFAULT_INPUT)
    p.add_argument('--pattern', default='*.input.pkl')
    p.add_argument('--cached', action='store_true')
    p.add_argument('--method', choices=['pte', 'reachable_te', 'dense_te'], default='pte')
    p.add_argument('--subset', help='Comma-separated commodity indices')
    p.add_argument('--suffix', default='')
    p.add_argument('--max-edges', type=int, default=5_000_000)
    p.add_argument('--max-states', type=int, default=2_500_000)
    p.add_argument('--preprocess-limit', type=float, default=180)
    p.add_argument('--solve-limit', type=float, default=180)
    args = p.parse_args()
    subset = [int(x) for x in args.subset.split(',')] if args.subset else None
    for path in sorted(args.input.glob(args.pattern)):
        run(path, args, args.cached or path.name.endswith('.input.pkl'), subset, args.method, args.suffix)
        gc.collect()


if __name__ == '__main__':
    main()
